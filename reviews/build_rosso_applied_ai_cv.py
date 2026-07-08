from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "ROSSO-HAN-APPLIED-AI-ENGINEERING-MANAGER-CV.docx"

INK = "172033"
BLUE = "1F4E78"
MUTED = "556070"
RULE = "B8C7D9"
BODY_SIZE = 10.2
BULLET_SIZE = 10.0
COMPACT_SIZE = 9.8
SECTION_SIZE = 10.5
ROLE_SIZE = 10.5
META_SIZE = 9.7


def set_font(run, size=None, bold=None, italic=None, color=INK, name="Arial"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def set_spacing(paragraph, before=0, after=0, line=1.0, keep_with_next=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep_with_next


def add_bottom_rule(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), RULE)
    p_bdr.append(bottom)


def configure_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    # Named resume override to compact_reference_guide.
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(BODY_SIZE)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.08

    for name in ["Title", "Subtitle", "Heading 1", "Heading 2"]:
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")

    bullet = styles["List Bullet"]
    bullet.font.name = "Arial"
    bullet.font.size = Pt(BULLET_SIZE)
    bullet.paragraph_format.left_indent = Inches(0.19)
    bullet.paragraph_format.first_line_indent = Inches(-0.13)
    bullet.paragraph_format.space_after = Pt(2.6)
    bullet.paragraph_format.line_spacing = 1.06

    if "Compact Skills" not in styles:
        compact = styles.add_style("Compact Skills", WD_STYLE_TYPE.PARAGRAPH)
    else:
        compact = styles["Compact Skills"]
    compact.base_style = styles["Normal"]
    compact.font.name = "Arial"
    compact.font.size = Pt(COMPACT_SIZE)
    compact.paragraph_format.space_after = Pt(2)
    compact.paragraph_format.line_spacing = 1.04


def add_header(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(p, after=1)
    set_font(p.add_run("ROSSO HAN"), size=22, bold=True, color=INK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(p, after=2)
    set_font(
        p.add_run("Engineering Manager | Applied AI, Agent Systems & RAG"),
        size=12,
        bold=True,
        color=BLUE,
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(p, after=5)
    set_font(
        p.add_run(
            "Oakland, CA | +1 (510) 458-9626 | jian.han3@gmail.com | "
            "linkedin.com/in/rossohan"
        ),
        size=9.2,
        color=MUTED,
    )


def add_section(doc, title, page_break_before=False):
    p = doc.add_paragraph()
    set_spacing(p, before=8, after=4, keep_with_next=True)
    p.paragraph_format.page_break_before = page_break_before
    set_font(p.add_run(title.upper()), size=SECTION_SIZE, bold=True, color=BLUE)
    add_bottom_rule(p)
    return p


def add_role(doc, title, company, dates, location=None):
    p = doc.add_paragraph()
    set_spacing(p, before=4.5, after=1.8, keep_with_next=True)
    set_font(p.add_run(title), size=ROLE_SIZE, bold=True, color=INK)
    set_font(p.add_run(f" | {company}"), size=ROLE_SIZE, bold=True, color=BLUE)
    set_font(p.add_run(f" | {dates}"), size=META_SIZE, color=MUTED)
    if location:
        set_font(p.add_run(f" | {location}"), size=META_SIZE, color=MUTED)


def add_bullet(doc, text, prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    set_spacing(p, after=2.6, line=1.06)
    if prefix and text.startswith(prefix):
        set_font(p.add_run(prefix), size=BULLET_SIZE, bold=True)
        set_font(p.add_run(text[len(prefix):]), size=BULLET_SIZE)
    else:
        set_font(p.add_run(text), size=BULLET_SIZE)
    return p


def add_summary(doc):
    add_section(doc, "Profile")
    p = doc.add_paragraph()
    set_spacing(p, after=3.5, line=1.08)
    set_font(
        p.add_run(
            "Engineering manager and hands-on Applied AI systems builder with 15+ years "
            "across backend platforms, cloud architecture, APIs, distributed systems, and "
            "technical leadership. Builds agentic AI and RAG products with memory, "
            "guardrails, evals, observability, human control, and reliable fallbacks. "
            "Combines production engineering judgment with founder-level execution from "
            "problem discovery through architecture, implementation, testing, and validation."
        ),
        size=BODY_SIZE,
    )


def add_core_strengths(doc):
    add_section(doc, "Core Strengths")
    p = doc.add_paragraph(style="Compact Skills")
    set_font(p.add_run("Applied AI: "), size=COMPACT_SIZE, bold=True, color=BLUE)
    set_font(
        p.add_run(
            "agentic workflows, RAG, LLM orchestration, AI memory, prompt systems, "
            "guardrails, evals, semantic search, vector databases"
        ),
        size=COMPACT_SIZE,
    )
    p = doc.add_paragraph(style="Compact Skills")
    set_font(p.add_run("Engineering Leadership: "), size=COMPACT_SIZE, bold=True, color=BLUE)
    set_font(
        p.add_run(
            "team development, architecture, roadmaps, execution, incident review, "
            "cross-functional alignment, distributed teams"
        ),
        size=COMPACT_SIZE,
    )
    p = doc.add_paragraph(style="Compact Skills")
    set_font(p.add_run("Platform Engineering: "), size=COMPACT_SIZE, bold=True, color=BLUE)
    set_font(
        p.add_run(
            "APIs, microservices, AWS, reliability, CI/CD, Terraform, observability, "
            "performance, production operations"
        ),
        size=COMPACT_SIZE,
    )


def add_ai_work(doc):
    add_section(doc, "Applied AI Work")
    add_role(
        doc,
        "Founder & Applied AI Systems Builder",
        "TryForgeAI",
        "2026 - Present",
        "Oakland, CA",
    )
    add_bullet(
        doc,
        "Forge AI RAG: Built a local mixed-format knowledge system that ingests Markdown, "
        "Excel, and Word files, creates Voyage AI embeddings, persists 487 indexed chunks "
        "in Chroma, retrieves relevant evidence, and generates grounded answers through Claude.",
        "Forge AI RAG:",
    )
    add_bullet(
        doc,
        "AI automation: Built a nine-node n8n workflow connecting scheduled execution, "
        "Google Sheets CRM data, lead filtering, Claude API reasoning, CRM updates, and "
        "Gmail notifications; created reusable skills for market intelligence and prospecting.",
        "AI automation:",
    )
    add_bullet(
        doc,
        "Avaloka AI: Implemented risk routing, LLM orchestration, guardian/repair/fallback, "
        "versioned prompts, and a SAGE-style memory writer, guardian, Care Card store, "
        "deterministic reader, lifecycle controls, diagnostics, and eval harness. "
        "Verified with 77 passing tests.",
        "Avaloka AI:",
    )
    add_bullet(
        doc,
        "Auralith: Built a modular AI-to-IoT pipeline that maps emotional context into "
        "guarded responses and safe low-stimulation light behavior, with AI shadow mode, "
        "scene validation, WLED adaptation, local-network controls, and a design-partner "
        "Pilot Kit. Verified with 200 passing tests across six modules.",
        "Auralith:",
    )
    add_bullet(
        doc,
        "Agent-first delivery: Created durable project operating systems that make product "
        "direction, architecture decisions, task state, traces, evals, and validation "
        "recoverable by humans and coding agents.",
        "Agent-first delivery:",
    )


def add_experience(doc):
    add_section(doc, "Professional Experience", page_break_before=True)
    add_role(doc, "Engineering Manager", "Rakuten Rewards", "2015 - Present")
    add_bullet(
        doc,
        "Led engineering teams responsible for backend and API infrastructure supporting "
        "a high-traffic consumer rewards platform."
    )
    add_bullet(
        doc,
        "Managed distributed collaboration across five time zones and aligned engineering, "
        "product, operations, and business stakeholders around delivery and reliability."
    )
    add_bullet(
        doc,
        "Guided architecture, planning, technical review, CI/CD practices, incident response, "
        "and continuous improvement for business-critical services."
    )
    add_bullet(
        doc,
        "Improved platform performance and resilience through caching, load balancing, "
        "microservice optimization, observability, and disciplined production operations."
    )
    add_bullet(
        doc,
        "Mentored engineers and translated ambiguous business and data-flow problems into "
        "executable technical plans."
    )

    add_role(doc, "Cloud Architect", "Hewlett-Packard (HP)", "2010 - 2015")
    add_bullet(
        doc,
        "Designed cloud-oriented APIs, service architecture, and resilient infrastructure "
        "for enterprise systems."
    )
    add_bullet(
        doc,
        "Supported modernization from monolithic applications toward scalable, maintainable "
        "services and more reliable data flows."
    )
    add_bullet(
        doc,
        "Introduced Infrastructure as Code practices with Terraform and improved repeatability "
        "across environment setup and deployment."
    )
    add_bullet(
        doc,
        "Worked across engineering and stakeholder groups to improve cloud readiness and "
        "operational reliability."
    )


def add_learning(doc):
    add_section(doc, "LLM & Enterprise RAG")
    p = doc.add_paragraph()
    set_spacing(p, after=3, line=1.06)
    set_font(
        p.add_run(
            "SupportVectors LLM and Enterprise RAG Bootcamp, Summer 2026. Applied study of "
            "embeddings, chunking, dense and sparse retrieval, fusion, reranking, query "
            "transformation, grounding, guardrails, RAG evaluation, RAPTOR, GraphRAG, "
            "semantic caching, SAGE precise retrieval, and Text-to-SQL. Course concepts are "
            "translated into Avaloka experiments, trace schemas, eval cases, and architecture "
            "decisions; advanced methods are treated as research until measured against a baseline."
        ),
        size=10.0,
    )


def add_toolbox(doc):
    add_section(doc, "Technical Toolbox")
    items = [
        (
            "AI / Agents",
            "OpenAI API, Anthropic Claude, LLM orchestration, RAG, Voyage AI, Chroma, "
            "semantic search, vector retrieval, memory systems, prompt registries, evals, guardrails",
        ),
        (
            "Automation",
            "n8n, agent skills, scheduled workflows, Google Sheets API, Gmail integrations, "
            "human-in-the-loop approval",
        ),
        (
            "Engineering",
            "Python, TypeScript, Java, React, Node.js, REST APIs, microservices, distributed systems",
        ),
        (
            "Cloud / Delivery",
            "AWS, Terraform, CI/CD, observability, incident management, caching, load balancing",
        ),
    ]
    for label, value in items:
        p = doc.add_paragraph(style="Compact Skills")
        set_font(p.add_run(f"{label}: "), size=COMPACT_SIZE, bold=True, color=BLUE)
        set_font(p.add_run(value), size=COMPACT_SIZE)


def add_education(doc):
    add_section(doc, "Education & Certifications")
    p = doc.add_paragraph(style="Compact Skills")
    set_font(
        p.add_run("M.S., E-Business Management"),
        size=COMPACT_SIZE,
        bold=True,
    )
    set_font(p.add_run(" - University of Warwick, UK"), size=COMPACT_SIZE)

    p = doc.add_paragraph(style="Compact Skills")
    set_font(p.add_run("B.S., Computer Science"), size=COMPACT_SIZE, bold=True)
    set_font(p.add_run(" - Beijing University of Technology, China"), size=COMPACT_SIZE)

    p = doc.add_paragraph(style="Compact Skills")
    set_font(p.add_run("Certifications: "), size=COMPACT_SIZE, bold=True)
    set_font(p.add_run("PMP; Certified ScrumMaster (CSM)"), size=COMPACT_SIZE)


def set_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_spacing(p)
    set_font(p.add_run("Rosso Han | Applied AI Engineering Manager"), size=7.5, color=MUTED)


def build():
    doc = Document()
    configure_styles(doc)
    set_footer(doc.sections[0])
    add_header(doc)
    add_summary(doc)
    add_core_strengths(doc)
    add_ai_work(doc)

    add_experience(doc)
    add_learning(doc)
    add_toolbox(doc)
    add_education(doc)

    properties = doc.core_properties
    properties.title = "Rosso Han - Applied AI Engineering Manager CV"
    properties.subject = "Applied AI, Agent Systems, RAG, Engineering Leadership"
    properties.author = "Rosso Han"
    properties.keywords = (
        "Applied AI, Engineering Manager, AI Agents, RAG, LLM, AI Evaluation, "
        "Guardrails, AI Memory, Platform Engineering"
    )
    doc.save(OUT)


if __name__ == "__main__":
    build()
